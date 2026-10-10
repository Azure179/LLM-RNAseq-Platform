"""Shared read-only annotation tables and enrichment gene/background contracts."""

from collections.abc import Iterator
from contextlib import contextmanager
import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from bulk_rnaseq.deg_data import REQUIRED_COLUMNS, read_deg_results
from bulk_rnaseq.exceptions import SkillError
from bulk_rnaseq.validation import MISSING, require_input_file


def validate_enrichment_parameters(pvalue: float, padj: float, timeout: float) -> None:
    for name, value in (("pvalue_threshold", pvalue), ("padj_threshold", padj), ("timeout", timeout)):
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise SkillError("INVALID_PARAMETER", f"{name} must be a finite number.")
    if not 0 <= pvalue <= 1 or not 0 <= padj <= 1 or timeout <= 0:
        raise SkillError("INVALID_PARAMETER", "Require thresholds in [0,1] and timeout > 0.")


@contextmanager
def table_rows(path: Path, sheet: str | None = None) -> Iterator[Iterator[list[Any]]]:
    """Stream UTF-8 CSV/TSV or read-only XLSX (preserving literal string IDs)."""
    require_input_file(path)
    try:
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            try:
                workbook = load_workbook(path, read_only=True, data_only=False)
            except (BadZipFile, InvalidFileException, ValueError) as cause:
                raise SkillError("MAPPING_FORMAT_ERROR", "Cannot read XLSX workbook.") from cause
            try:
                if sheet is None and len(workbook.sheetnames) != 1:
                    raise SkillError("MAPPING_SHEET_REQUIRED", "Choose mapping_sheet for a multi-sheet workbook.",
                                     details={"sheets": workbook.sheetnames})
                name = sheet or workbook.sheetnames[0]
                if name not in workbook.sheetnames:
                    raise SkillError("MAPPING_SHEET_NOT_FOUND", "Requested worksheet does not exist.")
                yield (list(row) for row in workbook[name].iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            if sheet is not None:
                raise SkillError("INVALID_PARAMETER", "mapping_sheet is only supported for XLSX.")
            with path.open(encoding="utf-8-sig", newline="") as handle:
                # Native .emapper.annotations is tab-separated, irrespective of suffix.
                delimiter = "," if path.suffix.lower() == ".csv" else "\t"
                yield csv.reader(handle, delimiter=delimiter, strict=True)
    except (UnicodeError, csv.Error, BadZipFile, ParseError) as cause:
        raise SkillError("MAPPING_FORMAT_ERROR", "Expected valid UTF-8 CSV/TSV or XLSX.") from cause
    except OSError as cause:
        raise SkillError("INPUT_READ_ERROR", "Cannot read annotation input file.", details={"path": str(path)}) from cause


def gene_identifier(value: Any, row: int) -> str:
    if (not isinstance(value, str) or value.strip().upper() in MISSING or value == "-"
            or value != value.strip() or any(char in value for char in "\r\n\t;")):
        raise SkillError("INVALID_IDENTIFIER", "Gene IDs must be non-missing text without surrounding whitespace, control characters or semicolons.",
                         details={"row": row})
    return value


def read_gene_list(path: Path, *, analysis_background: bool = False) -> tuple[frozenset[str], int]:
    """One ID per line or a gene_id column. Full DEG background excludes unselectable rows."""
    if path.suffix.lower() == ".xlsx":
        raise SkillError("GENE_LIST_FORMAT_ERROR", "Gene lists/background must be text or CSV/TSV; XLSX is mapping-only.")
    genes: set[str] = set()
    dropped = 0
    with table_rows(path) as iterator:
        rows = iter(iterator)
        first = next(rows, None)
        if not first:
            raise SkillError("GENE_LIST_EMPTY", "Gene list contains no genes.")
        if analysis_background and set(REQUIRED_COLUMNS) <= set(first):
            if path.suffix.lower() != ".csv":
                raise SkillError("GENE_LIST_FORMAT_ERROR", "Full DEG background must be the DEG Skill CSV.")
            records = read_deg_results(path)
            genes = {r.original["gene_id"] for r in records if r.padj is not None and r.log2fc is not None}
            dropped = len(records) - len(genes)
        else:
            if "gene_id" in first:
                if len(first) != len(set(first)):
                    raise SkillError("GENE_LIST_FORMAT_ERROR", "Gene list columns must be unique.")
                index, width = first.index("gene_id"), len(first)
            elif len(first) == 1:
                index, width = 0, 1
                genes.add(gene_identifier(first[0], 1))
            else:
                raise SkillError("GENE_LIST_FORMAT_ERROR", "Use one ID per line or a table with gene_id header.")
            for number, row in enumerate(rows, 2):
                if not row:
                    continue
                if len(row) != width:
                    raise SkillError("GENE_LIST_FORMAT_ERROR", "Gene list row width differs from header.", details={"row": number})
                gene = gene_identifier(row[index], number)
                dropped += gene in genes
                genes.add(gene)
    if not genes:
        raise SkillError("GENE_LIST_EMPTY", "No eligible genes remain in the gene list/background.")
    return frozenset(genes), dropped


@dataclass(frozen=True)
class EnrichmentInputs:
    query: frozenset[str]
    universe: frozenset[str]
    term_genes: dict[str, frozenset[str]]
    metrics: dict[str, Any]
    warnings: list[dict[str, Any]]


def prepare_enrichment(query: frozenset[str], background: frozenset[str],
                       pairs: frozenset[tuple[str, str]]) -> EnrichmentInputs:
    outside = query - background
    if outside:
        raise SkillError("GENES_OUTSIDE_UNIVERSE", "All target genes must belong to the declared eligible background.",
                         details={"count": len(outside), "examples": sorted(outside)[:10]})
    mapped = {gene for gene, _ in pairs}
    universe = background.intersection(mapped)
    foreground = query.intersection(mapped)
    if not foreground:
        raise SkillError("NO_MAPPED_GENES", "Gene list has no mappings in the selected annotation.")
    grouped: dict[str, set[str]] = {}
    for gene, term in pairs:
        if gene in universe:
            grouped.setdefault(term, set()).add(gene)
    metrics = {"n_query_input": len(query), "n_query_mapped": len(foreground),
               "n_background_input": len(background), "n_universe": len(universe),
               "n_query_unmapped": len(query - mapped), "n_background_unmapped": len(background - mapped),
               "n_terms_tested": len(grouped)}
    warnings: list[dict[str, Any]] = []
    for key in ("n_query_unmapped", "n_background_unmapped"):
        if metrics[key]:
            warnings.append({"code": "UNMAPPED_GENES", "message": f"{key}: genes excluded from effective analysis.", "count": metrics[key]})
    return EnrichmentInputs(foreground, universe, {term: frozenset(genes) for term, genes in sorted(grouped.items())}, metrics, warnings)
