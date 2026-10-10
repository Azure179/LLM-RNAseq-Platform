"""Dataset adapters only: preserve source data and reuse package validators/readers."""

from collections import Counter
from collections.abc import Iterable
import csv
from pathlib import Path
from typing import Any

from bulk_rnaseq.enrichment_inputs import gene_identifier, table_rows
from bulk_rnaseq.exceptions import SkillError
from bulk_rnaseq.go_mapping import read_go_mapping
from bulk_rnaseq.kegg_mapping import normalize_pathway, read_kegg_mapping
from bulk_rnaseq.run_io import sha256
from bulk_rnaseq.validation import require_input_file, validate_expression_inputs, validate_inputs, validate_pca_inputs


def local_path(root: Path, value: str) -> Path:
    path = Path(value)
    if path.is_absolute():
        raise SkillError("CONFIG_ERROR", "Use paths relative to Real_DATA_TEST.")
    result = (root / path).resolve()
    if not result.is_relative_to(root.resolve()):
        raise SkillError("CONFIG_ERROR", "Dataset paths must stay within Real_DATA_TEST.")
    return result


def write_table(path: Path, header: list[str], rows: Iterable, delimiter: str = ",") -> None:
    """Never overwrite a prepared input, even when rerunning with a different config."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise SkillError("PREPARED_INPUT_EXISTS", "Prepared input already exists; use --reuse-prepared after verifying provenance.", details={"path": str(path)})
    with path.open("x", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow(header)
        writer.writerows(rows)


def read_metadata_source(path: Path, has_header: bool, group_column: str) -> list[list[str]]:
    require_input_file(path)
    with path.open("rb") as handle:
        signature = handle.read(4)
    if signature.startswith(b"PK"):
        from openpyxl import load_workbook
        with path.open("rb") as handle:
            workbook = load_workbook(handle, read_only=True, data_only=False)
            try:
                if len(workbook.sheetnames) != 1:
                    raise SkillError("METADATA_FORMAT_ERROR", "Expected one metadata worksheet.")
                rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
            finally:
                workbook.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle, strict=True))
    if has_header:
        if not rows:
            raise SkillError("METADATA_FORMAT_ERROR", "Metadata is empty.")
        header = rows.pop(0)
        if len(header) != 2 or header[0] not in {"sample", "sample_id"} or header[1] != group_column:
            raise SkillError("METADATA_FORMAT_ERROR", "Expected sample,group metadata header.")
    for number, row in enumerate(rows, 1):
        if len(row) != 2:
            raise SkillError("METADATA_FORMAT_ERROR", "Metadata must contain exactly two columns.")
        for value in row:
            gene_identifier(value, number)
    return rows


def prepare_counts(root: Path, config: dict[str, Any]) -> dict[str, Any]:
    source = local_path(root, config["source_counts"])
    require_input_file(source)
    with source.open("rb") as handle:
        if handle.read(8).startswith((b"PK", b"\xd0\xcf\x11\xe0")):
            raise SkillError("COUNT_FORMAT_ERROR", "This dataset adapter expects TSV raw counts, not a binary Excel workbook.")
    with source.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t", strict=True)
        header = next(reader, None)
        rows = list(reader)
    if not header or header[0] not in {"gene_id", "Geneid"}:
        raise SkillError("COUNT_FORMAT_ERROR", "Expected gene_id/Geneid as first column.")
    header[0] = "gene_id"  # Column-label adapter only. All gene/sample IDs stay unchanged.
    if any(len(row) != len(header) for row in rows):
        raise SkillError("COUNT_FORMAT_ERROR", "Count row width differs from header.")
    metadata = read_metadata_source(local_path(root, config["source_metadata"]), config["source_metadata_has_header"], config["group_column"])
    prepared_dir = local_path(root, config.get("prepared_dir", "input"))
    all_counts, all_meta = prepared_dir / "count_matrix.all_samples.csv", prepared_dir / "metadata.all_samples.csv"
    write_table(all_counts, header, rows)
    write_table(all_meta, ["sample_id", config["group_column"]], metadata)
    overview = validate_expression_inputs(all_counts, all_meta, config["group_column"])
    if config["sample_scope"] != "comparison_only":
        raise SkillError("CONFIG_ERROR", "This runner uses the specified two-group comparison only.")
    wanted = {config["treatment"], config["control"]}
    selected_samples = {sample for sample, group in metadata if group in wanted}
    indexes = [0] + [i for i, name in enumerate(header[1:], 1) if name in selected_samples]
    selected_counts = local_path(root, config["count_matrix"])
    selected_meta = local_path(root, config["metadata"])
    write_table(selected_counts, [header[i] for i in indexes], ([row[i] for i in indexes] for row in rows))
    write_table(selected_meta, ["sample_id", config["group_column"], "condition"],
                ([sample, group, group] for sample, group in metadata if sample in selected_samples))
    checked = validate_inputs(selected_counts, selected_meta, config["treatment"], config["control"])
    validate_pca_inputs(selected_counts, selected_meta, config["group_column"])
    return {"n_genes": checked.gene_count, "n_samples": len(checked.sample_ids), "sample_names": list(checked.sample_ids),
            "groups": sorted(wanted), "replicates_per_group": dict(Counter(group for sample, group in metadata if sample in selected_samples)),
            "source_n_samples": len(overview.sample_ids), "source_groups": dict(Counter(overview.sample_groups)),
            "excluded_samples": sorted(set(overview.sample_ids) - selected_samples),
            "nonnegative_integer_counts": True, "sample_sets_match": True,
            "n_nonzero_genes": len(checked.analyzed_gene_ids), "genes": [row[0] for row in rows]}


def read_ontology(path: Path) -> tuple[dict[str, tuple[str, str]], str]:
    """Read GO namespace/name facts, not graph propagation or enrichment statistics."""
    if not path.is_file():
        return {}, "unavailable"
    lookup: dict[str, tuple[str, str]] = {}
    version = "unrecorded"
    record: dict[str, str] = {}
    namespace = {"biological_process": "BP", "molecular_function": "MF", "cellular_component": "CC"}

    def flush() -> None:
        if "id" in record and record.get("namespace") in namespace and record.get("is_obsolete") != "true":
            lookup[record["id"]] = (namespace[record["namespace"]], record.get("name", ""))

    with path.open(encoding="utf-8") as handle:
        for line in handle:
            line = line.rstrip("\n")
            if line.startswith("data-version: "):
                version = line.split(": ", 1)[1]
            if line.startswith("["):
                flush()
                record = {}
            elif ": " in line:
                key, value = line.split(": ", 1)
                record[key] = value
        flush()
    if not lookup:
        raise SkillError("ONTOLOGY_FORMAT_ERROR", "No active GO namespaces could be read from OBO.")
    return lookup, version


def prepare_annotations(root: Path, config: dict[str, Any], genes: set[str]) -> dict[str, Any]:
    source = local_path(root, config["source_annotations"])
    require_input_file(source)
    require_input_file(local_path(root, config["pathway_names"]))
    # Explicit crosswalk only; its provenance may be a reference or a user-requested naming rule.
    crosswalk: dict[str, set[str]] | None = None
    if config.get("gene_id_mapping"):
        crosswalk = {}
        with table_rows(local_path(root, config["gene_id_mapping"])) as rows:
            if next(rows, None) != ["gene_id", "query"]:
                raise SkillError("ID_MAPPING_FORMAT_ERROR", "Crosswalk must have gene_id,query TSV columns.")
            for number, row in enumerate(rows, 2):
                if len(row) != 2:
                    raise SkillError("ID_MAPPING_FORMAT_ERROR", "Expected two crosswalk columns.")
                gene, query = (gene_identifier(value, number) for value in row)
                if gene not in genes:
                    continue
                crosswalk.setdefault(query, set()).add(gene)
                if len(crosswalk[query]) > 1:
                    raise SkillError("AMBIGUOUS_ID_MAPPING", "One transcript/query maps to multiple counted genes.")
    def targets(query: str) -> set[str]:
        return crosswalk.get(query, set()) if crosswalk is not None else {query}

    ontology, version = read_ontology(local_path(root, config["go"]["ontology_file"]))
    go = read_go_mapping(source)
    all_go_genes = {gene for query, _ in go.pairs for gene in targets(query)}
    unknown = sorted(term for term in go.terms if term not in ontology)
    # Missing/obsolete GO IDs remain in a separate audit; never guessed or replaced.
    prepared_dir = local_path(root, config.get("prepared_dir", "input"))
    write_table(prepared_dir / "go_unresolved_terms.tsv", ["go_id", "reason"],
                ((term, "absent_or_obsolete_in_ontology" if ontology else "ontology_unavailable") for term in unknown), "\t")
    def go_rows() -> Iterable:
        for query, term in sorted(go.pairs):
            for gene in sorted(targets(query)):
                if ontology and term not in ontology:
                    continue
                group, name = ontology.get(term, ("", ""))
                yield [gene, term, group, name]
    write_table(local_path(root, config["gene_to_go"]), ["gene_id", "go_id", "ontology", "go_name"], go_rows(), "\t")
    go_usable = {gene for query, term in go.pairs for gene in targets(query)
                 if ontology.get(term, ("", ""))[0] == config["go"]["ontology"] or config["go"]["ontology"] == "ALL"}
    del go
    kegg = read_kegg_mapping(source)
    names: dict[str, str] = {}
    with local_path(root, config["pathway_names"]).open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle, delimiter="\t", strict=True):
            if len(row) != 2:
                raise SkillError("PATHWAY_NAME_FORMAT_ERROR", "Expected pathway ID and name in the name table.")
            canonical, _ = normalize_pathway(row[0])
            if canonical in names and names[canonical] != row[1]:
                raise SkillError("MAPPING_CONFLICT", "Conflicting pathway names.")
            names[canonical] = row[1]
    kegg_genes = {gene for query, _ in kegg.pairs for gene in targets(query)}
    write_table(local_path(root, config["gene_to_kegg"]), ["gene_id", "pathway_id", "pathway_name"],
                ([gene, term, names.get(term, "")] for query, term in sorted(kegg.pairs) for gene in sorted(targets(query))), "\t")
    write_table(prepared_dir / "pathway_id_aliases.tsv", ["input_pathway_id", "canonical_pathway_id"], sorted(kegg.aliases), "\t")
    return {"ontology_version": version, "ontology_sha256": sha256(local_path(root, config["go"]["ontology_file"])) if ontology else None,
            "unresolved_go_ids": len(unknown), "id_mapping_mode": "explicit_crosswalk" if crosswalk is not None else "exact_only",
            "id_mapping_basis": config.get("gene_id_mapping_basis", "user_supplied_unspecified") if crosswalk is not None else "exact_only",
            "go_any_annotation_genes": sorted(all_go_genes & genes), "go_usable_genes": sorted(go_usable & genes),
            "kegg_genes": sorted(kegg_genes & genes), "source_go_gene_coverage": len(all_go_genes & genes),
            "source_kegg_gene_coverage": len(kegg_genes & genes),
            "kegg_named_pathways": sum(term in names for term in kegg.terms), "kegg_pathways": len(kegg.terms)}
